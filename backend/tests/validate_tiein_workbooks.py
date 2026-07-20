"""Run the production TIE IN engine against two ILI Template workbooks.

Usage:
    python backend/tests/validate_tiein_workbooks.py OLD.xlsx NEW.xlsx
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


ENGINE_DIR = Path(__file__).resolve().parents[1] / "src" / "python_engines"
sys.path.insert(0, str(ENGINE_DIR))

import tiein_engine


HEADERS = (
    "feature_id",
    "log_distance_m",
    "joint_number",
    "dist_to_us_gw_m",
    "joint_length_m",
    "feature_type",
    "feature_identification",
    "dim_classification",
    "clock_position",
    "wall_thickness_mm",
    "length_mm",
    "width_mm",
    "depth_pct",
    "surface_location",
    "erf_asme",
    "psafe_bar",
    "comment",
)


def load_features(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "ILI Template" not in workbook.sheetnames:
            raise ValueError(f"{path}: worksheet 'ILI Template' was not found")
        sheet = workbook["ILI Template"]
        features = []
        blank_streak = 0
        for row in sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
            if all(value is None or str(value).strip() == "" for value in row):
                blank_streak += 1
                if blank_streak >= 200:
                    break
                continue
            blank_streak = 0
            feature = dict(zip(HEADERS, row))
            if feature["log_distance_m"] is None or feature["feature_type"] is None:
                continue
            features.append(feature)
        return features
    finally:
        workbook.close()


def summarize(old_path, new_path):
    old_features = load_features(old_path)
    new_features = load_features(new_path)
    result = tiein_engine.calculate_tiein_offset(old_features, new_features)
    transform = result.get("transform") or {}
    return {
        "success": result.get("success", False),
        "methodVersion": result.get("methodVersion"),
        "featureCount": {"old": len(old_features), "new": len(new_features)},
        "anchorCount": {
            "old": len(tiein_engine.extract_anchors(old_features, "OLD", [])),
            "new": len(tiein_engine.extract_anchors(new_features, "NEW", [])),
        },
        "decision": result.get("decision"),
        "tieIn": result.get("tieIn"),
        "transformDomain": transform.get("domain"),
        "evidenceOverlap": result.get("evidenceOverlap"),
        "knotCount": transform.get("knotCount", len(transform.get("knots") or [])),
        "alignedPairCount": result.get("alignedPairCount"),
        "spanResidual": result.get("spanResidual"),
        "robustFit": result.get("robustFit"),
        "error": result.get("error"),
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Expected OLD.xlsx and NEW.xlsx")
    summary = summarize(Path(sys.argv[1]), Path(sys.argv[2]))
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    if not summary["success"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
