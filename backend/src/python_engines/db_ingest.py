import argparse
import csv
import hashlib
import io
import json
import re
import sys
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


CURRENT_DIR = Path(__file__).resolve().parent
SRC_DIR = CURRENT_DIR.parent
BACKEND_DIR = SRC_DIR.parent
PROJECT_ROOT = BACKEND_DIR.parent
for module_path in (PROJECT_ROOT, SRC_DIR):
    if str(module_path) not in sys.path:
        sys.path.insert(0, str(module_path))

import engine.constants as constants
from backend.engine.parser_dict import (
    auto_detect_field_map,
    load_feature_type_normalization,
    load_field_synonym_lookup,
    normalize_field_name,
)
from engine.constants import UNIVERSAL_SCHEMA
from engine.loader import find_best_col, normalize_clock, normalize_feature_type, normalize_surface
from python_engines.models import ILIFeature, Pipeline, SessionLocal, SurveyRun, init_db


DATABASE_FIELDS = [
    field_name for field_name, *_ in UNIVERSAL_SCHEMA if not field_name.startswith("_")
]
NUMERIC_FIELDS = {
    "log_distance_m",
    "joint_number",
    "dist_to_us_gw_m",
    "joint_length_m",
    "wall_thickness_mm",
    "length_mm",
    "width_mm",
    "depth_pct",
    "erf_asme",
    "psafe_bar",
}
FLOAT_FIELDS = NUMERIC_FIELDS - {"joint_number"}
INSERT_BATCH_SIZE = 5_000
TRAILING_EMPTY_ROW_LIMIT = 256
NO_ILI_HEADER_MESSAGE = "No worksheet contains a recognisable ILI header row."
XLSX_MAGIC = b"PK"
LEGACY_XLS_MAGIC = bytes.fromhex("D0CF11E0")
PARSER_VERSION = "agnostic-ili-parser-2026.07"
SCHEMA_VERSION = "universal-ili-v1"
QUALITY_GATE_VERSION = "ingestion-quality-gate-2026.07"

# A unit only counts when it is an explicit token: parenthesised "(ft)",
# bracketed "[in]", dotted "ft." / "in.", or an unambiguous full word.  The
# bare English words "in" ("Depth in % WT") and "ft" must never trigger a
# conversion.
_IMPERIAL_DISTANCE_TOKEN_RE = re.compile(
    r"[\(\[]\s*(?:ft|feet|foot)\s*\.?\s*[\)\]]"  # "(ft)", "[ft]", "(feet.)"
    r"|\bft\.(?!\w)"                             # "ft." with an explicit dot
    r"|\b(?:feet|foot)\b",                       # unambiguous full words
    re.IGNORECASE,
)
_IMPERIAL_DIMENSION_TOKEN_RE = re.compile(
    r"[\(\[]\s*(?:in|inch|inches)\s*\.?\s*[\)\]]"  # "(in)", "[in]", "(inch.)"
    r"|\bin\.(?!\w)"                               # "in." with an explicit dot
    r"|\b(?:inch|inches)\b",                       # unambiguous full words
    re.IGNORECASE,
)
_METRIC_DISTANCE_TOKEN_RE = re.compile(
    r"[\(\[]\s*m\s*[\)\]]|\b(?:meters?|metres?)\b", re.IGNORECASE,
)
_METRIC_DIMENSION_TOKEN_RE = re.compile(
    r"[\(\[]\s*mm\s*[\)\]]|\bmm\b|\bmillimet(?:er|re)s?\b", re.IGNORECASE,
)

DISTANCE_FIELDS = {"log_distance_m", "dist_to_us_gw_m", "joint_length_m"}
DIMENSION_FIELDS = {"wall_thickness_mm", "length_mm", "width_mm"}

FT_TO_M = 0.3048
IN_TO_MM = 25.4


def detect_source_units(field_sources):
    """Infer the unit of each mapped distance/dimension column from its own header.

    ``field_sources`` maps a universal field name to the source header(s)
    actually bound to it.  Scanning every header joined into one string used to
    let the word "in" in an unrelated column ("Depth in % WT") convert
    wall/length/width by 25.4, and a stray "ft" rescale the whole odometer
    axis.  A unit is now assigned per column, only from an explicit token, and
    an explicit metric marker ("(mm)", "(m)", "mm") in the same header wins
    over an imperial one.

    Returns ``{field_name: "imperial"|"metric"|None}``.
    """
    units = {}
    for field, headers in field_sources.items():
        if field in DISTANCE_FIELDS:
            imperial_re, metric_re = _IMPERIAL_DISTANCE_TOKEN_RE, _METRIC_DISTANCE_TOKEN_RE
        elif field in DIMENSION_FIELDS:
            imperial_re, metric_re = _IMPERIAL_DIMENSION_TOKEN_RE, _METRIC_DIMENSION_TOKEN_RE
        else:
            continue
        unit = None
        for header in headers:
            text = str(header)
            if metric_re.search(text):
                unit = "metric"
                break
            if imperial_re.search(text):
                unit = "imperial"
        units[field] = unit
    return units


def _summarize_source_units(field_units, fields):
    """Collapse per-field units into the audit's legacy group value."""
    group_units = {field_units.get(field) for field in fields}
    if "imperial" in group_units:
        return "imperial"
    if "metric" in group_units:
        return "metric"
    return None


def build_quality_gate(audit):
    """Return a conservative intake decision without approving engineering use."""
    review_reasons = []
    if int(audit.get("accepted_rows", 0)) == 0:
        return {
            "version": QUALITY_GATE_VERSION,
            "status": "BLOCKED",
            "requires_engineering_review": True,
            "review_reasons": ["no_accepted_rows"],
        }

    if int(audit.get("rejected_rows", 0)) > 0:
        review_reasons.append("rejected_source_rows")
    review_reasons.extend(
        f"numeric_coercion:{field}"
        for field, count in audit.get("numeric_coercions", {}).items()
        if int(count) > 0
    )
    review_reasons.extend(
        field
        for field, count in audit.get("quality_warnings", {}).items()
        if int(count) > 0
    )

    return {
        "version": QUALITY_GATE_VERSION,
        "status": "REVIEW_REQUIRED" if review_reasons else "READY_FOR_ENGINEERING_REVIEW",
        "requires_engineering_review": True,
        "review_reasons": review_reasons,
    }


def source_sha256(file_path, chunk_size=1024 * 1024):
    digest = hashlib.sha256()
    with Path(file_path).open("rb") as source:
        for chunk in iter(lambda: source.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _source_kind(file_path):
    path = Path(file_path)
    with path.open("rb") as source:
        magic = source.read(4)

    if magic.startswith(XLSX_MAGIC):
        return "excel"
    if magic.startswith(LEGACY_XLS_MAGIC):
        raise ValueError(
            "The legacy .xls format is not supported; re-save the file as .xlsx or .csv."
        )
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        return "excel"
    return "csv"


def detect_header_row(preview, synonym_lookup, max_scan=30):
    best_index = None
    best_fields = set()

    for index, row in enumerate(preview.head(max_scan).itertuples(index=False, name=None)):
        matched_fields = {
            synonym_lookup.get(normalize_field_name(value))
            for value in row
            if pd.notna(value)
        }
        matched_fields.discard(None)
        if "log_distance_m" in matched_fields and len(matched_fields) > len(best_fields):
            best_index = index
            best_fields = matched_fields

    if best_index is None or len(best_fields) < 2:
        raise ValueError(
            "No header row found: log_distance_m and at least one other ILI column "
            "must be recognisable."
        )

    return best_index, len(best_fields)


def _row_has_value(row):
    return any(value is not None and str(value).strip() != "" for value in row)


def _unique_headers(header_row):
    headers = []
    occurrences = {}
    for index, value in enumerate(header_row):
        base = str(value).strip() if value is not None and str(value).strip() else f"_col{index}"
        occurrence = occurrences.get(base, 0)
        occurrences[base] = occurrence + 1
        headers.append(base if occurrence == 0 else f"{base}.{occurrence}")
    return headers


def _tabulate_sheet(rows, synonym_lookup, preview_rows):
    """Turn one worksheet's rows into a frame, or None when it holds no header.

    Shared by both readers below so the fast path and the fallback cannot drift
    apart on what counts as a header row, a blank row, or a short row.
    """
    preview = rows[:preview_rows]
    if not preview:
        return None
    try:
        header_index, matched_fields = detect_header_row(
            pd.DataFrame(preview),
            synonym_lookup,
            preview_rows,
        )
    except ValueError:
        return None

    headers = _unique_headers(rows[header_index])
    width = len(headers)
    data_rows = []
    trailing_empty_rows = 0

    for row in rows[header_index + 1:]:
        values = tuple(row[:width])
        if _row_has_value(values):
            trailing_empty_rows = 0
            # A row whose trailing cells were never written arrives short, and
            # pandas rejects a ragged list outright — one such row aborted the
            # whole ingestion. Rows are padded to the header width instead,
            # which is what the missing cells mean.
            data_rows.append(
                values if len(values) == width else values + (None,) * (width - len(values))
            )
        else:
            trailing_empty_rows += 1
            if trailing_empty_rows >= TRAILING_EMPTY_ROW_LIMIT:
                break

    return matched_fields, pd.DataFrame(data_rows, columns=headers), header_index


def _calamine_cell(value):
    """Translate one calamine value into what openpyxl would have returned.

    Two differences matter downstream. An unset cell arrives as "" rather than
    None, and the importer reads None — ``dropna``, ``combine_first`` and the
    blank-row audit all change meaning against an empty string. A whole number
    arrives as float, so a vendor's integer feature number would reach a text
    column as "7001.0".
    """
    if value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _read_sheets_calamine(file_path):
    """Yield each worksheet's rows using the Rust reader, or None when absent.

    openpyxl spends most of an ILI import inside its own XML layer: on a
    120k-row workbook it costs about 17 s where calamine costs under 2 s, and
    that single stage was three quarters of the wall clock an operator waits
    through on the intake screen. The dependency is optional so an environment
    that has not installed it yet still imports, only slowly.
    """
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        return None

    workbook = CalamineWorkbook.from_path(str(file_path))

    def sheets():
        for sheet_name in workbook.sheet_names:
            sheet = workbook.get_sheet_by_name(sheet_name)
            # skip_empty_area=False keeps leading blank rows, so a header sits
            # at the same index openpyxl would report it at.
            rows = sheet.to_python(skip_empty_area=False)
            yield [tuple(_calamine_cell(cell) for cell in row) for row in rows]

    return sheets()


def _read_sheets_openpyxl(file_path):
    """Yield each worksheet's rows with openpyxl, one sheet at a time."""
    source = Path(file_path).open("rb")
    try:
        workbook = load_workbook(
            source,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception:
        source.close()
        raise
    try:
        for worksheet in workbook.worksheets:
            yield [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
        source.close()


def _select_ili_sheet(sheets, synonym_lookup, preview_rows):
    selected = None
    for rows in sheets:
        tabulated = _tabulate_sheet(rows, synonym_lookup, preview_rows)
        if tabulated is None:
            continue
        # Only the best sheet is retained; the rest are released as we go, so a
        # multi-sheet workbook costs one sheet of memory rather than all of them.
        if selected is None or tabulated[0] > selected[0]:
            selected = tabulated

    if selected is None:
        raise ValueError(NO_ILI_HEADER_MESSAGE)

    matched_fields, frame, header_index = selected
    return frame, header_index, matched_fields


def _load_xlsx_table(file_path, synonym_lookup, preview_rows):
    sheets = _read_sheets_calamine(file_path)
    if sheets is not None:
        try:
            return _select_ili_sheet(sheets, synonym_lookup, preview_rows)
        except ValueError:
            # The workbook was read and holds no ILI header. openpyxl reads the
            # same cells, so retrying only doubles the wait before the operator
            # sees the same message.
            raise
        except Exception as error:
            # A workbook the fast reader cannot decode must still import. The
            # fallback is the reader every file was accepted by until now, so
            # the file is retried rather than rejected — recorded on stderr so
            # a file that keeps needing the slow path can be found.
            print(
                json.dumps({
                    "event": "xlsx_fast_reader_fallback",
                    "file": Path(file_path).name,
                    "error": f"{type(error).__name__}: {error}",
                }),
                file=sys.stderr,
            )
    return _select_ili_sheet(
        _read_sheets_openpyxl(file_path),
        synonym_lookup,
        preview_rows,
    )


def load_source_table(file_path, synonym_lookup, preview_rows=30):
    source_kind = _source_kind(file_path)
    if source_kind == "excel":
        raw, header_index, matched_fields = _load_xlsx_table(
            file_path,
            synonym_lookup,
            preview_rows,
        )
    else:
        preview = pd.read_csv(file_path, header=None, nrows=preview_rows)
        header_index, matched_fields = detect_header_row(preview, synonym_lookup, preview_rows)
        raw = pd.read_csv(file_path, header=header_index)

    # Blank rows are dropped before mapping, but the audit's source-row count is
    # a conservation claim: measuring it after this filter under-reported the
    # file and made "accepted + rejected == source" true by construction.
    rows_before_blank_filter = int(len(raw))
    raw = raw.dropna(how="all").reset_index(drop=True)
    raw.attrs["source_rows"] = rows_before_blank_filter
    raw.attrs["blank_rows"] = rows_before_blank_filter - int(len(raw))
    raw.columns = [
        str(column).strip() if pd.notna(column) else f"_col{index}"
        for index, column in enumerate(raw.columns)
    ]
    if raw.empty:
        raise ValueError("The ILI file has no data rows after its header.")
    return raw, header_index, matched_fields


def _nonempty_mask(series):
    return series.notna() & ~series.astype(str).str.strip().str.casefold().isin(
        {"", "nan", "none", "null", "n/a", "-"}
    )


def summarize_feature_counts(frame):
    """Build the wizard's small intake summary without exposing source rows."""
    feature_types = frame.get("feature_type", pd.Series(dtype="object"))
    labels = feature_types.fillna("").astype(str).str.casefold()
    categories = pd.Series("other", index=frame.index, dtype="object")

    # Keep the precedence consistent with the legacy browser parser: a metal
    # loss near a weld remains an anomaly, never an anchor.
    # NOTE: these must be single-backslash \b word boundaries.  Written as \\b
    # inside a raw string the pattern matches a *literal* backslash-b, so every
    # POF short code (ANOM, ML, COCL, GW, VALV, OFFT) silently fell through to
    # "other" and the wizard reported zero anomalies and zero welds.
    categories.loc[labels.str.contains(r"metal loss|corrosion|\banom\b|\bcocl\b|\bcorr\b|\bml\b", regex=True)] = "anomaly"
    categories.loc[(categories == "other") & labels.str.contains(r"weld|girth|seam|\bgw\b", regex=True)] = "weld"
    categories.loc[(categories == "other") & labels.str.contains(r"valve|\bvalv\b", regex=True)] = "valve"
    categories.loc[(categories == "other") & labels.str.contains(r"offtake|off take|\bofft\b", regex=True)] = "offtake"
    categories.loc[(categories == "other") & labels.str.contains(r"tee", regex=True)] = "tee"
    categories.loc[(categories == "other") & labels.str.contains(r"bend|elbow", regex=True)] = "bend"

    counts = categories.value_counts()
    welds = int(counts.get("weld", 0))
    valves = int(counts.get("valve", 0))
    tees = int(counts.get("tee", 0))
    offtakes = int(counts.get("offtake", 0))
    return {
        "total": int(len(frame)),
        "anomalies": int(counts.get("anomaly", 0)),
        "welds": welds,
        "valves": valves,
        "tees": tees,
        "offtakes": offtakes,
        "bends": int(counts.get("bend", 0)),
        "others": int(counts.get("other", 0)),
        # Offtakes are counted here because the matcher counts them: DTW's
        # _alignment_feature_kind normalises "offtake" to "tee", so they are
        # Tier-1 anchors. Excluding them made the wizard promise fewer anchors
        # than the engine would actually use.
        "gw_anchors": welds + valves + tees + offtakes,
        # Invalid distances are rejected before a canonical run is persisted.
        "invalid_anchor_distances": 0,
    }


def normalize_ingestion_frame(raw, field_map):
    if raw.empty:
        raise ValueError("The ILI file has no data rows.")

    universal = pd.DataFrame(
        {field: pd.Series(pd.NA, index=raw.index, dtype="object") for field in DATABASE_FIELDS}
    )

    # Track which source header(s) fed each universal column, so unit
    # detection can inspect the mapped column's own header only.
    bound_sources = {}
    for raw_column, universal_column in field_map.items():
        if universal_column not in DATABASE_FIELDS or raw_column not in raw.columns:
            continue
        if universal[universal_column].isna().all():
            universal[universal_column] = raw[raw_column].copy()
        else:
            universal[universal_column] = universal[universal_column].combine_first(raw[raw_column])
        bound_sources.setdefault(universal_column, []).append(raw_column)

    for universal_column in DATABASE_FIELDS:
        if _nonempty_mask(universal[universal_column]).any():
            continue
        matched_column, similarity = find_best_col(universal_column, raw.columns, 0.55)
        if matched_column is not None and similarity >= 0.55:
            universal[universal_column] = raw[matched_column]
            # The fallback replaces the column wholesale, so it also replaces
            # any previously recorded (empty) source header.
            bound_sources[universal_column] = [matched_column]

    if not _nonempty_mask(universal["log_distance_m"]).any():
        raise ValueError("The log_distance_m column could not be mapped.")

    numeric_coercions = {}
    for field in NUMERIC_FIELDS:
        source_has_value = _nonempty_mask(universal[field])
        converted = pd.to_numeric(universal[field], errors="coerce")
        finite = converted.notna() & np.isfinite(converted)

        if field == "joint_number":
            integer = finite & np.isclose(converted.fillna(0), np.round(converted.fillna(0)))
            finite &= integer

        coerced_count = int((source_has_value & ~finite).sum())
        if coerced_count:
            numeric_coercions[field] = coerced_count
        universal[field] = converted.where(finite, np.nan)

    field_units = detect_source_units(bound_sources)
    unit_conversions = {}
    for field in DISTANCE_FIELDS:
        if field_units.get(field) != "imperial":
            continue
        mask = universal[field].notna()
        if mask.any():
            universal.loc[mask, field] = universal.loc[mask, field] * FT_TO_M
            unit_conversions[field] = "ft_to_m"
    for field in DIMENSION_FIELDS:
        if field_units.get(field) != "imperial":
            continue
        mask = universal[field].notna()
        if mask.any():
            universal.loc[mask, field] = universal.loc[mask, field] * IN_TO_MM
            unit_conversions[field] = "in_to_mm"
    source_units = {
        "distances": _summarize_source_units(field_units, DISTANCE_FIELDS),
        "dimensions": _summarize_source_units(field_units, DIMENSION_FIELDS),
    }

    universal["clock_position"] = universal["clock_position"].apply(
        lambda value: None if pd.isna(value) else normalize_clock(value)
    )
    universal["surface_location"] = universal["surface_location"].apply(
        lambda value: None if pd.isna(value) else normalize_surface(value)
    )
    feature_types = universal["feature_type"].values
    feature_identifications = universal["feature_identification"].values
    normalized_types = [
        normalize_feature_type(
            None if pd.isna(feature_types[index]) else feature_types[index],
            None if pd.isna(feature_identifications[index]) else feature_identifications[index],
        )
        for index in range(len(universal))
    ]
    # Normalisation collapses a vendor label onto a canonical class, and that
    # loses information downstream consumers need. "Cluster corrosion" becomes
    # "Metal Loss", and the alignment engine's cluster rule — which reads
    # feature_type, feature_identification and dim_classification — then had no
    # column left carrying the word "cluster", so vendor-merged groups were fed
    # into anomaly matching after all. The source label is kept in
    # feature_identification whenever the vendor left that column empty.
    universal["feature_identification"] = [
        feature_identifications[index]
        if not pd.isna(feature_identifications[index])
        and str(feature_identifications[index]).strip() != ""
        else (
            None if pd.isna(feature_types[index])
            else str(feature_types[index]).strip() or None
        )
        for index in range(len(universal))
    ]
    universal["feature_type"] = normalized_types

    valid_distance = universal["log_distance_m"].notna()
    rejected_rows = int((~valid_distance).sum())
    universal = universal.loc[valid_distance].copy()
    if universal.empty:
        raise ValueError("No row has a numeric, finite log_distance_m.")

    universal["joint_number"] = universal["joint_number"].apply(
        lambda value: int(value) if pd.notna(value) else None
    )
    universal = universal.sort_values("log_distance_m", kind="stable").reset_index(drop=True)

    quality_warnings = {}
    # A depth outside 0-100 %WT is not a measurement this schema can carry: it
    # would reach B31G as a depth ratio above 1 (or below 0) and produce an
    # invalid remaining-strength result presented as an engineering number. The
    # value is cleared and counted; the row itself is conserved, because its
    # position and identity remain valid evidence.
    invalid_depth = universal["depth_pct"].notna() & ~universal["depth_pct"].between(0, 100)
    invalid_depth_count = int(invalid_depth.sum())
    if invalid_depth_count:
        quality_warnings["depth_pct_out_of_range"] = invalid_depth_count
        universal.loc[invalid_depth, "depth_pct"] = np.nan

    blank_rows = int(raw.attrs.get("blank_rows", 0))
    source_rows = int(raw.attrs.get("source_rows", len(raw)))
    rejection_reasons = {}
    if rejected_rows:
        rejection_reasons["invalid_log_distance_m"] = rejected_rows
    if blank_rows:
        rejection_reasons["blank_row"] = blank_rows

    audit = {
        "source_rows": source_rows,
        "accepted_rows": int(len(universal)),
        "rejected_rows": rejected_rows + blank_rows,
        "rejection_reasons": rejection_reasons,
        "numeric_coercions": numeric_coercions,
        "unit_conversions": unit_conversions,
        "source_units": source_units,
        "quality_warnings": quality_warnings,
        "feature_counts": summarize_feature_counts(universal),
    }
    audit["quality_gate"] = build_quality_gate(audit)
    return universal, audit


def _database_value(field, value):
    if value is None or pd.isna(value):
        return None
    if field == "joint_number":
        return int(value)
    if field in FLOAT_FIELDS:
        return float(value)
    if field == "feature_id" and isinstance(value, (float, np.floating)) and value.is_integer():
        return str(int(value))
    cleaned = str(value).strip()
    return cleaned or None


def _insert_feature_batches(session, run_id, frame, batch_size=INSERT_BATCH_SIZE):
    if session.bind is not None and session.bind.dialect.name == "postgresql":
        _copy_insert_features(session, run_id, frame, batch_size=batch_size)
        return

    for start in range(0, len(frame), batch_size):
        chunk = frame.iloc[start : start + batch_size]
        mappings = []
        for record in chunk.to_dict(orient="records"):
            mapping = {"run_id": run_id}
            mapping.update(
                {field: _database_value(field, record.get(field)) for field in DATABASE_FIELDS}
            )
            mappings.append(mapping)
        session.bulk_insert_mappings(ILIFeature, mappings, return_defaults=False)


def _copy_insert_features(session, run_id, frame, batch_size=INSERT_BATCH_SIZE):
    """Bulk-load normalized rows with PostgreSQL COPY, retaining a test fallback."""
    raw_connection = session.connection().connection
    copy_columns = ["run_id", *DATABASE_FIELDS]
    copy_sql = (
        f"COPY ili_features ({', '.join(copy_columns)}) "
        "FROM STDIN WITH (FORMAT CSV, NULL '\\N')"
    )
    cursor = raw_connection.cursor()
    try:
        for start in range(0, len(frame), batch_size):
            chunk = frame.iloc[start : start + batch_size]
            buffer = io.StringIO()
            writer = csv.writer(buffer, lineterminator="\n")
            for values in chunk.itertuples(index=False, name=None):
                record = dict(zip(DATABASE_FIELDS, values))
                row_values = [run_id]
                for field in DATABASE_FIELDS:
                    value = _database_value(field, record.get(field))
                    row_values.append(r"\N" if value is None else value)
                writer.writerow(row_values)
            buffer.seek(0)
            cursor.copy_expert(copy_sql, buffer)
    finally:
        cursor.close()


def parse_survey_date(value):
    """An ISO calendar date, or None when the caller did not supply one.

    An unparseable value is rejected rather than dropped: the run would be
    stored with a blank date that reads as "this run predates the date field",
    which is a different fact from "the operator typed something we ignored".
    """
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError as error:
        failure = ValueError(f"Survey date is not a valid ISO calendar date: {text}")
        failure.code = "INVALID_SURVEY_DATE"
        raise failure from error


def ingest_csv(
    file_path,
    pipeline_name,
    survey_year,
    vendor,
    organization_id,
    user_id=None,
    source_name=None,
    source_object_key=None,
    survey_date=None,
):
    synonym_lookup = load_field_synonym_lookup()
    if not synonym_lookup:
        error = RuntimeError(
            "The parser synonym dictionary is empty; run seed_dictionary.py first."
        )
        error.code = "PARSER_DICTIONARY_UNAVAILABLE"
        raise error

    feature_type_map = load_feature_type_normalization()
    if feature_type_map:
        constants.FEATURE_TYPE_NORMALIZATION.update(feature_type_map)

    raw, header_index, matched_fields = load_source_table(file_path, synonym_lookup)
    field_map = auto_detect_field_map(raw.columns, synonym_lookup)
    normalized, audit = normalize_ingestion_frame(raw, field_map)
    audit.update(
        {
            "header_row": header_index + 1,
            "recognized_header_fields": matched_fields,
            "mapped_columns": len(field_map),
        }
    )

    checksum = source_sha256(file_path)
    surveyed_on = parse_survey_date(survey_date)
    session = SessionLocal()
    try:
        pipeline = session.query(Pipeline).filter(
            Pipeline.name == pipeline_name,
            Pipeline.organization_id == organization_id,
        ).first()
        if pipeline is None:
            pipeline = Pipeline(name=pipeline_name, organization_id=organization_id)
            session.add(pipeline)
            session.flush()

        run = SurveyRun(
            pipeline_id=pipeline.id,
            # The date wins where both are present: the caller derives the year
            # from it, so they cannot disagree, and a run flown in December is
            # not the same evidence as one flown the following January.
            survey_year=surveyed_on.year if surveyed_on else survey_year,
            survey_date=surveyed_on,
            vendor=vendor,
            source_file=source_name or Path(file_path).name,
            source_object_key=source_object_key,
            source_sha256=checksum,
            uploaded_by=user_id,
            row_count=len(normalized),
            parser_version=PARSER_VERSION,
            schema_version=SCHEMA_VERSION,
            ingestion_audit=audit,
        )
        session.add(run)
        session.flush()

        _insert_feature_batches(session, run.id, normalized)
        session.commit()
        run_id = run.id
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()

    return {
        "status": "success",
        "run_id": run_id,
        "row_count": len(normalized),
        "summary": audit["feature_counts"],
        "audit": audit,
        "parser_version": PARSER_VERSION,
        "schema_version": SCHEMA_VERSION,
        "source_sha256": checksum,
        "survey_date": surveyed_on.isoformat() if surveyed_on else None,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--pipeline", required=True)
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--vendor", required=True)
    parser.add_argument("--organization-id", type=int, required=True)
    parser.add_argument("--user-id", type=int, required=True)
    parser.add_argument("--source-name")
    parser.add_argument("--source-object-key")
    parser.add_argument("--survey-date")
    args = parser.parse_args()

    try:
        init_db()
        result = ingest_csv(
            args.file,
            args.pipeline,
            args.year,
            args.vendor,
            args.organization_id,
            user_id=args.user_id,
            source_name=args.source_name,
            source_object_key=args.source_object_key,
            survey_date=args.survey_date,
        )
        print("JSON_OUTPUT:" + json.dumps(result))
        return 0
    except Exception as error:
        payload = {"status": "error", "message": str(error)}
        error_code = getattr(error, "code", None)
        if error_code:
            payload["code"] = error_code
        print("JSON_OUTPUT:" + json.dumps(payload))
        return 1


if __name__ == "__main__":
    sys.exit(main())
